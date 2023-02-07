import openpyxl
from openpyxl.styles import Border,Side,Font,PatternFill,Alignment
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
wb = Workbook()
sheet = wb.active
chart = BarChart()

sales_data = [
              ["Product","Online","Store"],
              [1,3,4],
              [2,4,5],
              [8,9,5],
              [4,9,2]
]

for row in sales_data:
    sheet.append(row)

data = Reference(sheet, min_col=2, max_col=3, min_row=1, max_row=8)
chart.add_data(data, titles_from_data=True)
sheet.add_chart(chart,"E2")

wb.save("C://Users/user773/Desktop/barchartdata.xlsx")